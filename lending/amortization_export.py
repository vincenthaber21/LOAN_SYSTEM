"""Excel and PDF exports of a stored declining-balance amortization schedule."""

from __future__ import annotations

import io
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

def _plain(value) -> str:
    text = "" if value is None else str(value)
    for src, dest in (
        ("\u2014", "-"),
        ("\u2013", "-"),
        ("\u00b7", " / "),
        ("\u2022", "-"),
        ("\u20b1", "PHP "),
    ):
        text = text.replace(src, dest)
    return text

HEADER_FILL = "163A14"
HEADER_FONT = "FFFFFF"
MONEY = '#,##0.00'


def _money(amount) -> str:
    value = Decimal(str(amount or 0)).quantize(Decimal("0.01"))
    return f"{value:,.2f}"


def _date(value) -> str:
    if not value:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _summary_lines(loan, schedule):
    borrower = loan.application.borrower_name
    return [
        (f"{loan.reference} amortization", f"{borrower}"),
        ("Principal", _money(schedule["rows"][0]["balance"] if schedule["rows"] else loan.principal)),
        ("Payment per period", _money(schedule["payment"])),
        ("Payments", str(schedule["number_of_payments"])),
        ("Interval (days)", str(schedule["interval_days"])),
        ("Total repayment", _money(schedule["total_repayment"])),
        ("Total interest", _money(schedule["total_interest"])),
        ("Add-on rate", f"{schedule['add_on_percent']}%"),
        ("Rate per period", f"{schedule['period_rate_percent']}%"),
        ("Nominal annual rate", f"{schedule['nominal_annual_percent']}%"),
        ("Effective annual rate", f"{schedule['effective_annual_percent']}%"),
        ("Comparison", schedule["comparison"]),
    ]


def build_amortization_xlsx(loan, schedule) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Amortization"
    bold = Font(bold=True, color=HEADER_FONT)
    fill = PatternFill("solid", fgColor=HEADER_FILL)
    thin = Border(
        left=Side(style="thin", color="D5E5D8"),
        right=Side(style="thin", color="D5E5D8"),
        top=Side(style="thin", color="D5E5D8"),
        bottom=Side(style="thin", color="D5E5D8"),
    )
    label_font = Font(bold=True, color="163A14")

    for index, (label, value) in enumerate(_summary_lines(loan, schedule), start=1):
        sheet.cell(index, 1, label).font = label_font
        sheet.cell(index, 2, value)

    header_row = 14
    headers = ["No.", "Due Date", "Scheduled Payment", "Interest", "Principal", "Balance"]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin

    for offset, row in enumerate(schedule["rows"]):
        excel_row = header_row + 1 + offset
        values = [
            row["number"],
            _date(row["due_date"]),
            float(row["payment"] or 0),
            float(row["interest"] or 0),
            float(row["principal"] or 0),
            float(row["balance"] or 0),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(excel_row, column, value)
            cell.border = thin
            if column >= 3:
                cell.number_format = MONEY
                cell.alignment = Alignment(horizontal="right")

    widths = [8, 16, 22, 16, 16, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.column_dimensions["B"].width = 88
    sheet.row_dimensions[12].height = 32
    sheet.cell(12, 2).alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_amortization_pdf(loan, schedule) -> bytes:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"{loan.reference} amortization",
    )
    title_style = ParagraphStyle(
        "AmortTitle",
        fontName="Helvetica-Bold",
        fontSize=14,
        textColor=colors.HexColor("#163a14"),
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "AmortBody",
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#163a14"),
    )
    story = [
        Paragraph(_plain(f"{loan.reference} — {loan.application.borrower_name}"), title_style),
        Paragraph(_plain(schedule["comparison"]), body_style),
        Spacer(1, 3 * mm),
        Paragraph(
            _plain(
                "Payment per period "
                f"{_money(schedule['payment'])} · "
                f"Total repayment {_money(schedule['total_repayment'])} · "
                f"Total interest {_money(schedule['total_interest'])} · "
                f"Rate per period {schedule['period_rate_percent']}% · "
                f"Nominal annual {schedule['nominal_annual_percent']}% · "
                f"EAR {schedule['effective_annual_percent']}%"
            ),
            body_style,
        ),
        Spacer(1, 4 * mm),
    ]
    header = ["No.", "Due Date", "Scheduled Payment", "Interest", "Principal", "Balance"]
    data = [header]
    for row in schedule["rows"]:
        data.append(
            [
                str(row["number"]),
                _date(row["due_date"]) or "—",
                _money(row["payment"]),
                _money(row["interest"]),
                _money(row["principal"]),
                _money(row["balance"]),
            ]
        )
    table = Table(data, repeatRows=1, colWidths=[18 * mm, 32 * mm, 42 * mm, 36 * mm, 36 * mm, 40 * mm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#163a14")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
                ("ALIGN", (0, 0), (1, -1), "CENTER"),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#e0f6eb")),
                ("ROWBACKGROUNDS", (0, 2), (-1, -1), [colors.white, colors.HexColor("#f4faf6")]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d5e5d8")),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 4 * mm))
    story.append(
        Paragraph(
            _plain(f"Generated {timezone.localtime().strftime('%b %d, %Y %I:%M %p')}"),
            body_style,
        )
    )
    document.build(story)
    return buffer.getvalue()

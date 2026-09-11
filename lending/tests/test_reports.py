from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from lending.models import Loan, LoanApplication, LoanProduct, Payment, User
from mutual_aid.models import MutualAidContribution, MutualAidMembership, MutualAidPlan
from savings.models import SavingsAccount, SavingsProduct, SavingsTransaction


class AuditReportExportTests(TestCase):
    def setUp(self):
        self.password = "secret-pass"
        self.manager = User.objects.create_user(
            username="report-manager",
            email="report-manager@example.com",
            password=self.password,
            full_name="Report Manager",
            role=User.Role.MANAGER,
        )
        self.member = User.objects.create_user(
            username="report-member",
            email="report-member@example.com",
            password=self.password,
            full_name="Report Member",
            role=User.Role.MEMBER,
        )
        self.product = LoanProduct.objects.create(
            name="Audit Personal",
            loan_type=LoanProduct.LoanType.PERSONAL,
            min_amount=Decimal("1000.00"),
            max_amount=Decimal("50000.00"),
            interest_rate=Decimal("2.50"),
        )
        self.application = LoanApplication.objects.create(
            borrower=self.member,
            loan_product=self.product,
            amount_requested=Decimal("10000.00"),
            term_months=6,
            status=LoanApplication.Status.DISBURSED,
            applied_on=date(2026, 8, 5),
            created_by=self.manager,
            form_ref_no="APP-AUDIT-001",
        )
        LoanApplication.objects.filter(pk=self.application.pk).update(
            created_at=timezone.make_aware(datetime(2026, 8, 5, 9, 15, 30)),
        )
        self.application.refresh_from_db()
        self.loan = Loan.objects.create(
            application=self.application,
            principal=Decimal("10000.00"),
            interest_rate=Decimal("2.50"),
            term_months=6,
            disbursed_date=date(2026, 8, 10),
            total_payable=Decimal("11500.00"),
            outstanding_balance=Decimal("11000.00"),
            disbursement_method="Cash",
            disbursement_reference="REL-001",
            disbursed_by=self.manager,
        )
        self.payment = Payment.objects.create(
            loan=self.loan,
            amount=Decimal("500.00"),
            savings_adjustment=Decimal("5.00"),
            mutual_aid_contribution=Decimal("15.00"),
            payment_date=date(2026, 8, 20),
            method=Payment.Method.CASH,
            reference_number="PAY-AUDIT-001",
            recorded_by=self.manager,
        )
        Payment.objects.create(
            loan=self.loan,
            amount=Decimal("200.00"),
            payment_date=date(2026, 7, 1),
            method=Payment.Method.CASH,
            reference_number="PAY-OUT-OF-RANGE",
            recorded_by=self.manager,
        )
        savings_product = SavingsProduct.objects.create(
            name="Regular Savings",
            interest_rate=Decimal("2.00"),
        )
        savings_account = SavingsAccount.objects.create(
            member=self.member,
            product=savings_product,
            balance=Decimal("1000.00"),
            opened_by=self.manager,
        )
        self.savings_tx = SavingsTransaction.objects.create(
            account=savings_account,
            transaction_type=SavingsTransaction.Type.DEPOSIT,
            amount=Decimal("250.00"),
            method=SavingsTransaction.Method.CASH,
            reference_number="SV-AUDIT-001",
            balance_after=Decimal("1000.00"),
            notes="Counter deposit",
            created_at=timezone.make_aware(datetime(2026, 8, 15, 14, 45, 10)),
            created_by=self.manager,
        )
        plan = MutualAidPlan.objects.create(
            name="KAP Aid",
            contribution_amount=Decimal("15.00"),
            max_benefit_amount=Decimal("10000.00"),
        )
        membership = MutualAidMembership.objects.create(
            member=self.member,
            plan=plan,
            enrolled_by=self.manager,
        )
        self.contribution = MutualAidContribution.objects.create(
            membership=membership,
            amount=Decimal("15.00"),
            method=MutualAidContribution.Method.CASH,
            reference_number="MA-AUDIT-001",
            notes="Weekly aid",
            created_at=timezone.make_aware(datetime(2026, 8, 18, 11, 5, 0)),
            recorded_by=self.manager,
        )
        self.client.force_login(self.manager)

    def _export(self, **params):
        query = {
            "period": "custom",
            "from": "2026-08-01",
            "to": "2026-09-11",
            "grain": "day",
            "staff": "",
            "product": "",
        }
        query.update(params)
        url_name = "export_cashflow_csv" if query.get("format") == "csv" else "export_audit_xlsx"
        return self.client.get(reverse(url_name), query)

    def _workbook(self, response):
        return load_workbook(BytesIO(response.content), data_only=False)

    def _sheet_values(self, worksheet):
        return [
            str(cell.value)
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        ]

    def test_excel_workbook_has_a_tab_for_each_feature(self):
        response = self._export()
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response["Content-Type"],
        )
        self.assertIn(
            "audit-report-day-all-staff-2026-08-01-to-2026-09-11.xlsx",
            response["Content-Disposition"],
        )
        workbook = self._workbook(response)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Cover",
                "All Activity",
                "Collections",
                "Applications",
                "Disbursements",
                "Savings",
                "Mutual Aid",
                "Period Totals",
            ],
        )
        generated = timezone.localtime()
        cover_values = self._sheet_values(workbook["Cover"])
        self.assertIn("Generated month", cover_values)
        self.assertIn(generated.strftime("%B %Y"), cover_values)
        self.assertIn("2026-08-01", cover_values)
        self.assertIn("2026-09-11", cover_values)
        self.assertIn("PAY-AUDIT-001", self._sheet_values(workbook["Collections"]))
        self.assertIn("APP-AUDIT-001", self._sheet_values(workbook["Applications"]))
        self.assertIn(self.loan.reference, self._sheet_values(workbook["Disbursements"]))
        self.assertIn("SV-AUDIT-001", self._sheet_values(workbook["Savings"]))
        self.assertIn("MA-AUDIT-001", self._sheet_values(workbook["Mutual Aid"]))
        self.assertNotIn("PAY-OUT-OF-RANGE", self._sheet_values(workbook["Collections"]))
        self.assertTrue(workbook["Collections"].auto_filter.ref)
        self.assertEqual(workbook["Collections"].freeze_panes, "A6")
        collections = workbook["Collections"]
        total_row = None
        for row in collections.iter_rows(min_row=6, max_col=1):
            if row[0].value == "Total":
                total_row = row[0].row
                break
        self.assertIsNotNone(total_row)
        # Amount column (I) should be a computed number, not an Excel formula.
        amount_total = collections.cell(total_row, 9).value
        self.assertEqual(amount_total, 500.0)
        self.assertFalse(isinstance(amount_total, str) and str(amount_total).startswith("="))

    def test_product_filter_limits_loan_rows(self):
        other_product = LoanProduct.objects.create(
            name="Other Product",
            loan_type=LoanProduct.LoanType.BUSINESS,
            min_amount=Decimal("1000.00"),
            max_amount=Decimal("80000.00"),
            interest_rate=Decimal("3.00"),
        )
        other_application = LoanApplication.objects.create(
            borrower=self.member,
            loan_product=other_product,
            amount_requested=Decimal("8000.00"),
            term_months=4,
            status=LoanApplication.Status.SUBMITTED,
            applied_on=date(2026, 8, 8),
            form_ref_no="APP-OTHER-001",
        )
        response = self._export(product=str(self.product.pk))
        workbook = self._workbook(response)
        application_values = self._sheet_values(workbook["Applications"])
        self.assertIn("APP-AUDIT-001", application_values)
        self.assertNotIn("APP-OTHER-001", application_values)
        self.assertEqual(other_application.product_name, "Other Product")

    def test_csv_fallback_still_available(self):
        response = self._export(format="csv")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        text = response.content.decode("utf-8")
        self.assertIn("PAY-AUDIT-001", text)
        self.assertIn("Detailed collections", text)

    def test_reports_page_links_to_period_audit_excel(self):
        response = self.client.get(
            reverse("reports"),
            {
                "period": "custom",
                "from": "2026-08-01",
                "to": "2026-09-11",
                "grain": "day",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Download Excel")
        self.assertContains(response, reverse("export_audit_xlsx"))
        self.assertContains(response, "from=2026-08-01")
        self.assertContains(response, "to=2026-09-11")
        self.assertContains(response, "tab for collections")

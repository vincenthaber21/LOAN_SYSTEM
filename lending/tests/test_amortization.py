"""Add-on rate converted to a declining-balance amortization schedule."""

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from lending.amortization import (
    MATCH_TOTAL_REPAYMENT,
    STANDARD_ROUNDING,
    AmortizationError,
    addon_rate_for_term,
    calculate_amortization,
    payment_count_for_term,
    solve_period_rate,
)
from lending.models import Loan, LoanApplication, LoanProduct, User
from lending.services import generate_schedule, rebuild_loan_schedule


class DecliningBalanceScheduleTests(TestCase):
    def test_example_addons_to_declining_balance(self):
        """P=450000, add-on 33%, 26 payments, every 14 days from 2026-09-25.

        The period rate is solved from
        A = P * i / (1 - (1 + i) ** -n). For the rounded payment 23,019.23 that
        rate is 0.022390017491 (2.2390% per period), which is the value that
        actually satisfies the equation.
        """
        result = calculate_amortization(
            Decimal("450000"),
            Decimal("0.33"),
            26,
            14,
            date(2026, 9, 25),
        )
        self.assertEqual(result.payment, Decimal("23019.23"))
        self.assertEqual(result.total_repayment, Decimal("598500.00"))
        self.assertEqual(result.total_interest, Decimal("148500.00"))
        self.assertEqual(result.period_rate, Decimal("0.022390017491"))
        self.assertEqual(result.period_rate_percent, Decimal("2.2390"))
        self.assertEqual(result.nominal_annual_percent, Decimal("58.37"))
        self.assertEqual(result.effective_annual_percent, Decimal("78.12"))

        solved = solve_period_rate(Decimal("450000"), Decimal("23019.23"), 26)
        self.assertAlmostEqual(float(solved), float(result.period_rate), places=10)

        self.assertEqual(result.rows[0].balance, Decimal("450000.00"))
        self.assertEqual(result.rows[0].payment, Decimal("0.00"))

        first = result.payments[0]
        self.assertEqual(first.due_date, date(2026, 9, 25))
        self.assertEqual(first.payment, Decimal("23019.23"))
        self.assertEqual(first.interest, Decimal("10075.51"))
        self.assertEqual(first.principal, Decimal("12943.72"))
        self.assertEqual(first.balance, Decimal("437056.28"))

        second = result.payments[1]
        self.assertEqual(second.due_date, date(2026, 10, 9))
        self.assertEqual(second.interest, Decimal("9785.70"))
        self.assertEqual(second.principal, Decimal("13233.53"))
        self.assertEqual(second.balance, Decimal("423822.75"))

        last = result.payments[-1]
        self.assertEqual(last.number, 26)
        # (26 - 1) * 14 days after 2026-09-25.
        self.assertEqual(last.due_date, date(2027, 9, 10))
        self.assertEqual(last.due_date, date(2026, 9, 25) + timedelta(days=25 * 14))
        self.assertEqual(last.balance, Decimal("0.00"))
        self.assertEqual(sum(row.payment for row in result.payments), Decimal("598500.00"))

        again = calculate_amortization(
            Decimal("450000"),
            Decimal("0.33"),
            26,
            14,
            date(2026, 9, 25),
            period_rate=result.period_rate,
        )
        self.assertEqual(again.payments[0], result.payments[0])
        self.assertEqual(again.payments[-1], result.payments[-1])

    def test_single_payment(self):
        result = calculate_amortization(
            Decimal("1000"),
            Decimal("0.05"),
            1,
            30,
            date(2026, 1, 1),
        )
        self.assertEqual(len(result.payments), 1)
        self.assertEqual(result.payment, Decimal("1050.00"))
        self.assertEqual(result.payments[0].principal, Decimal("1000.00"))
        self.assertEqual(result.payments[0].interest, Decimal("50.00"))
        self.assertEqual(result.payments[0].payment, Decimal("1050.00"))
        self.assertEqual(result.payments[0].balance, Decimal("0.00"))
        self.assertEqual(result.payments[0].due_date, date(2026, 1, 1))

    def test_zero_addon_rate(self):
        result = calculate_amortization(
            Decimal("1000"),
            Decimal("0"),
            3,
            7,
            date(2026, 2, 2),
            final_payment_policy=MATCH_TOTAL_REPAYMENT,
        )
        self.assertEqual(result.period_rate, Decimal("0"))
        self.assertEqual(result.payment, Decimal("333.33"))
        self.assertEqual([row.interest for row in result.payments], [Decimal("0.00")] * 3)
        self.assertEqual([row.payment for row in result.payments], [
            Decimal("333.33"),
            Decimal("333.33"),
            Decimal("333.34"),
        ])
        self.assertEqual(result.payments[-1].balance, Decimal("0.00"))
        self.assertEqual(sum(row.payment for row in result.payments), Decimal("1000.00"))
        self.assertEqual(result.payments[-1].due_date, date(2026, 2, 2) + timedelta(days=14))

    def test_weekly_and_monthly_intervals(self):
        start = date(2026, 4, 1)
        for interval in (7, 30):
            result = calculate_amortization(
                Decimal("20000"),
                Decimal("0.12"),
                10,
                interval,
                start,
            )
            self.assertEqual(len(result.payments), 10)
            self.assertGreater(result.period_rate, 0)
            self.assertEqual(result.payments[-1].balance, Decimal("0.00"))
            self.assertEqual(
                result.payments[-1].due_date,
                start + timedelta(days=9 * interval),
            )
            self.assertEqual(sum(row.payment for row in result.payments), result.total_repayment)
            for previous, current in zip(result.payments, result.payments[1:]):
                self.assertEqual((current.due_date - previous.due_date).days, interval)

    def test_standard_rounding_still_ends_at_zero(self):
        result = calculate_amortization(
            Decimal("450000"),
            Decimal("0.33"),
            26,
            14,
            date(2026, 9, 25),
            final_payment_policy=STANDARD_ROUNDING,
        )
        self.assertEqual(result.payments[-1].balance, Decimal("0.00"))
        self.assertEqual(
            result.payments[-1].payment,
            result.payments[-1].principal + result.payments[-1].interest,
        )

    def test_input_validation(self):
        start = date(2026, 1, 1)
        with self.assertRaises(AmortizationError):
            calculate_amortization(Decimal("0"), Decimal("0.1"), 12, 30, start)
        with self.assertRaises(AmortizationError):
            calculate_amortization(Decimal("1000"), Decimal("-0.01"), 12, 30, start)
        with self.assertRaises(AmortizationError):
            calculate_amortization(Decimal("1000"), Decimal("0.1"), 0, 30, start)
        with self.assertRaises(AmortizationError):
            calculate_amortization(Decimal("1000"), Decimal("0.1"), 12, 0, start)
        with self.assertRaises(AmortizationError):
            calculate_amortization(Decimal("1000"), Decimal("0.1"), 12, 30, "2026-01-01")

    def test_term_mapping_for_twelve_month_biweekly_loan(self):
        self.assertEqual(payment_count_for_term(12, 14), 26)
        self.assertEqual(payment_count_for_term(1, 30), 1)
        self.assertEqual(addon_rate_for_term(Decimal("2.75"), 12), Decimal("0.33"))


class LoanScheduleStorageTests(TestCase):
    def setUp(self):
        self.officer = User.objects.create_user(
            username="amort-officer",
            email="amort-officer@example.com",
            password="secret-pass",
            full_name="Amort Officer",
            role=User.Role.OFFICER,
        )
        self.member = User.objects.create_user(
            username="amort-member",
            email="amort-member@example.com",
            password="secret-pass",
            full_name="Amort Member",
            role=User.Role.MEMBER,
        )
        self.product = LoanProduct.objects.create(
            name="Biweekly Capital",
            loan_type=LoanProduct.LoanType.PERSONAL,
            min_amount=Decimal("1000.00"),
            max_amount=Decimal("500000.00"),
            interest_rate=Decimal("2.75"),
        )
        self.application = LoanApplication.objects.create(
            borrower=self.member,
            loan_product=self.product,
            amount_requested=Decimal("450000.00"),
            term_months=12,
            payment_frequency=LoanApplication.PaymentFrequency.BIWEEKLY,
            final_interest_rate=Decimal("2.75"),
            final_term_months=12,
            status=LoanApplication.Status.ACTIVE,
        )
        self.loan = Loan.objects.create(
            application=self.application,
            principal=Decimal("450000.00"),
            disbursed_principal=Decimal("450000.00"),
            interest_rate=Decimal("2.75"),
            term_months=12,
            disbursed_date=date(2026, 9, 25),
            total_payable=Decimal("0.00"),
            outstanding_balance=Decimal("0.00"),
            grace_period_days=0,
        )

    def test_schedule_is_stored_and_reused(self):
        self.assertTrue(rebuild_loan_schedule(self.loan))
        self.loan.refresh_from_db()
        self.assertEqual(self.loan.period_rate, Decimal("0.022390017491"))
        self.assertEqual(self.loan.addon_rate, Decimal("0.330000"))
        self.assertEqual(self.loan.payment_interval_days, 14)
        self.assertEqual(self.loan.number_of_payments, 26)
        self.assertEqual(self.loan.total_payable, Decimal("598500.00"))
        self.assertEqual(self.loan.outstanding_balance, Decimal("598500.00"))

        first = self.loan.installments.get(installment_number=1)
        self.assertEqual(first.due_date, date(2026, 9, 25))
        self.assertEqual(first.interest_component, Decimal("10075.51"))
        self.assertEqual(first.principal_component, Decimal("12943.72"))
        self.assertEqual(first.amount_due, Decimal("23019.23"))
        self.assertEqual(first.ending_balance, Decimal("437056.28"))

        last = self.loan.installments.get(installment_number=26)
        self.assertEqual(last.due_date, date(2027, 9, 10))
        self.assertEqual(last.ending_balance, Decimal("0.00"))

        stored_rate = self.loan.period_rate
        generate_schedule(self.loan)
        self.loan.refresh_from_db()
        self.assertEqual(self.loan.period_rate, stored_rate)
        self.assertEqual(
            self.loan.installments.get(installment_number=2).interest_component,
            Decimal("9785.70"),
        )

    def test_officer_schedule_page_and_exports(self):
        self.client.force_login(self.officer)
        page = self.client.get(reverse("officer_repayment_schedule", args=[self.loan.pk]))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "₱23,019.23")
        self.assertContains(page, "₱10,075.51")
        self.assertContains(page, "₱437,056.28")
        self.assertContains(page, "Sep 10, 2027")
        self.assertContains(page, "2.2390%")
        self.assertContains(page, "58.37%")
        self.assertContains(page, "78.12%")

        detail = self.client.get(reverse("officer_loan_detail", args=[self.loan.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "2.2390%")
        self.assertContains(detail, "₱23,019.23")

        workbook_response = self.client.get(reverse("export_schedule_xlsx", args=[self.loan.pk]))
        self.assertEqual(workbook_response.status_code, 200)
        self.assertIn("spreadsheetml", workbook_response["Content-Type"])
        workbook = load_workbook(BytesIO(workbook_response.content))
        sheet = workbook.active
        self.assertEqual(sheet.cell(14, 1).value, "No.")
        self.assertEqual(sheet.cell(14, 3).value, "Scheduled Payment")
        self.assertAlmostEqual(sheet.cell(16, 3).value, 23019.23, places=2)
        self.assertAlmostEqual(sheet.cell(16, 4).value, 10075.51, places=2)

        pdf_response = self.client.get(reverse("export_schedule_amortization_pdf", args=[self.loan.pk]))
        self.assertEqual(pdf_response.status_code, 200)
        self.assertEqual(pdf_response["Content-Type"], "application/pdf")
        self.assertTrue(pdf_response.content.startswith(b"%PDF"))

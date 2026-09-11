from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from django.utils import timezone

from .cashflow_reports import audit_detail_context
from .models import Features

THEME = "186010"
THEME_DEEP = "0F2E0C"
MINT = "E7F2E4"
ALT_ROW = "F7FBF6"
GOLD = "B8860B"
LINE = "C9D6C4"
WHITE = "FFFFFF"
MUTED = "5B6B57"

CURRENCY = '"₱"#,##0.00'
DATE_FMT = "YYYY-MM-DD"
THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)

LEDGER_COLUMNS = [
    ("Category", "category", "text", 16),
    ("Date", "iso_date", "date", 12),
    ("Month", "month", "text", 14),
    ("Time", "time", "text", 13),
    ("Reference", "reference", "text", 18),
    ("Member", "member", "text", 22),
    ("Detail", "detail", "text", 28),
    ("Type", "type", "text", 16),
    ("Amount", "amount", "money", 14),
    ("Method / status", "method_or_status", "text", 16),
    ("Recorded by", "recorded_by", "text", 18),
    ("Notes", "notes", "text", 36),
]

COLLECTION_COLUMNS = [
    ("Date", "iso_date", "date", 12),
    ("Month", "month", "text", 14),
    ("Time", "time", "text", 13),
    ("Payment reference", "reference", "text", 18),
    ("Loan", "loan", "text", 14),
    ("Application", "application", "text", 16),
    ("Member", "member", "text", 22),
    ("Product", "product", "text", 18),
    ("Amount", "amount", "money", 14),
    ("Applied to loan", "loan_applied", "money", 16),
    ("Savings adjustment", "savings_adjustment", "money", 18),
    ("Mutual aid", "mutual_aid", "money", 14),
    ("Method", "method", "text", 14),
    ("Installment", "installment", "text", 12),
    ("Recorded by", "recorded_by", "text", 18),
]

APPLICATION_COLUMNS = [
    ("Date", "iso_date", "date", 12),
    ("Month", "month", "text", 14),
    ("Time", "time", "text", 13),
    ("Reference", "reference", "text", 18),
    ("Member", "member", "text", 22),
    ("Email", "email", "text", 24),
    ("Product", "product", "text", 18),
    ("Amount requested", "amount", "money", 16),
    ("Term (months)", "term_months", "int", 14),
    ("Status", "status", "text", 14),
    ("Created by", "created_by", "text", 18),
    ("Reviewed by", "reviewed_by", "text", 18),
]

DISBURSEMENT_COLUMNS = [
    ("Date", "iso_date", "date", 12),
    ("Month", "month", "text", 14),
    ("Time", "time", "text", 13),
    ("Loan", "reference", "text", 14),
    ("Receipt", "receipt", "text", 14),
    ("Application", "application", "text", 16),
    ("Member", "member", "text", 22),
    ("Product", "product", "text", 18),
    ("Principal", "principal", "money", 14),
    ("Net release", "net_release", "money", 14),
    ("Outstanding", "outstanding", "money", 14),
    ("Status", "status", "text", 12),
    ("Method", "method", "text", 16),
    ("Disbursement reference", "disbursement_reference", "text", 20),
    ("Disbursed by", "disbursed_by", "text", 18),
]

SAVINGS_COLUMNS = [
    ("Date", "iso_date", "date", 12),
    ("Month", "month", "text", 14),
    ("Time", "time", "text", 13),
    ("Reference", "reference", "text", 18),
    ("Account", "account", "text", 16),
    ("Member", "member", "text", 22),
    ("Product", "product", "text", 18),
    ("Type", "type", "text", 14),
    ("Amount", "amount", "money", 14),
    ("Method", "method", "text", 16),
    ("Balance after", "balance_after", "money", 14),
    ("Recorded by", "recorded_by", "text", 18),
    ("Notes", "notes", "text", 32),
]

MUTUAL_AID_COLUMNS = [
    ("Date", "iso_date", "date", 12),
    ("Month", "month", "text", 14),
    ("Time", "time", "text", 13),
    ("Reference", "reference", "text", 18),
    ("Membership", "membership", "text", 16),
    ("Member", "member", "text", 22),
    ("Plan", "plan", "text", 18),
    ("Period", "period", "text", 22),
    ("Amount", "amount", "money", 14),
    ("Method", "method", "text", 16),
    ("Recorded by", "recorded_by", "text", 18),
    ("Notes", "notes", "text", 32),
]


def _fill(color):
    return PatternFill("solid", fgColor=color)


def _font(*, size=11, bold=False, color="1A1A1A", name="Calibri"):
    return Font(name=name, size=size, bold=bold, color=color)


def _as_number(value):
    if value in ("", None):
        return None
    if isinstance(value, Decimal):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return value


def _as_date(value):
    if not value:
        return None
    if hasattr(value, "year") and not isinstance(value, datetime):
        return value
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return value


def _cell_value(kind, row, key):
    value = row.get(key, "")
    if kind == "date":
        return _as_date(row.get("iso_date") or value)
    if kind == "money":
        return _as_number(value)
    if kind == "int":
        if value in ("", None):
            return None
        return int(value)
    return value if value is not None else ""


def _style_header_row(ws, row_idx, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row_idx, col)
        cell.fill = _fill(THEME)
        cell.font = _font(size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    ws.row_dimensions[row_idx].height = 22


def _apply_print(ws, title, org_name):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.leftMargin = 0.4
    ws.page_setup.rightMargin = 0.4
    ws.page_setup.topMargin = 0.6
    ws.page_setup.bottomMargin = 0.6
    ws.oddHeader.left.text = org_name
    ws.oddHeader.right.text = "CONFIDENTIAL"
    ws.oddFooter.left.text = title
    ws.oddFooter.right.text = "Page &P of &N"
    ws.sheet_view.showGridLines = False


def _write_banner(ws, org_name, title, subtitle, col_count, tab_color=THEME):
    ws.sheet_properties.tabColor = tab_color
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=col_count)
    brand = ws.cell(1, 1, org_name)
    brand.fill = _fill(THEME_DEEP)
    brand.font = _font(size=12, bold=True, color=WHITE)
    brand.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22
    heading = ws.cell(2, 1, title)
    heading.fill = _fill(THEME)
    heading.font = _font(size=16, bold=True, color=WHITE)
    heading.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 26
    sub = ws.cell(3, 1, subtitle)
    sub.fill = _fill(MINT)
    sub.font = _font(size=10, color=THEME_DEEP)
    sub.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 18
    fills = {
        1: _fill(THEME_DEEP),
        2: _fill(THEME),
        3: _fill(MINT),
    }
    for row, fill in fills.items():
        for col in range(1, col_count + 1):
            ws.cell(row, col).fill = fill


def _write_table(ws, columns, rows, header_row=5):
    col_count = len(columns)
    for index, (title, _key, kind, width) in enumerate(columns, start=1):
        ws.cell(header_row, index, title)
        ws.column_dimensions[get_column_letter(index)].width = width
    _style_header_row(ws, header_row, col_count)

    if not rows:
        ws.merge_cells(
            start_row=header_row + 1,
            start_column=1,
            end_row=header_row + 1,
            end_column=col_count,
        )
        empty = ws.cell(header_row + 1, 1, "No records in this reporting period.")
        empty.font = _font(size=10, color=MUTED)
        empty.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[header_row + 1].height = 22
        return header_row, header_row

    money_cols = []
    money_totals = {}
    last_row = header_row
    for offset, row in enumerate(rows, start=1):
        last_row = header_row + offset
        stripe = ALT_ROW if offset % 2 == 0 else WHITE
        for col_idx, (title, key, kind, _width) in enumerate(columns, start=1):
            value = _cell_value(kind, row, key)
            cell = ws.cell(last_row, col_idx, value)
            cell.border = THIN
            cell.fill = _fill(stripe)
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=kind == "text" and title in {"Notes", "Detail"},
                horizontal="right" if kind in {"money", "int"} else "left",
            )
            if kind == "money":
                cell.number_format = CURRENCY
                money_cols.append(col_idx)
                amount = value if isinstance(value, (int, float)) else 0
                money_totals[col_idx] = money_totals.get(col_idx, 0) + (amount or 0)
            elif kind == "date":
                cell.number_format = DATE_FMT
            elif kind == "int":
                cell.number_format = "0"

    unique_money = sorted(set(money_cols))
    if unique_money:
        total_row = last_row + 1
        for col_idx in range(1, col_count + 1):
            cell = ws.cell(total_row, col_idx)
            cell.fill = _fill(MINT)
            cell.border = THIN
            cell.font = _font(size=10, bold=True, color=THEME_DEEP)
            if col_idx == 1:
                cell.value = "Total"
                cell.font = _font(size=11, bold=True, color=THEME_DEEP)
            elif col_idx in unique_money:
                cell.value = round(money_totals.get(col_idx, 0), 2)
                cell.number_format = CURRENCY
                cell.alignment = Alignment(horizontal="right", vertical="center")
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(col_count)}{last_row}"
        return header_row, total_row

    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(col_count)}{last_row}"
    return header_row, last_row


def _add_data_sheet(wb, name, title, subtitle, org_name, columns, rows, tab_color=THEME):
    ws = wb.create_sheet(name)
    _write_banner(ws, org_name, title, subtitle, len(columns), tab_color=tab_color)
    header_row, _end = _write_table(ws, columns, rows)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.print_title_rows = f"1:{header_row}"
    _apply_print(ws, title, org_name)
    return ws


def _meta_pairs(cashflow, details, generated_local, generated_by, org_name):
    return [
        ("Organization", org_name),
        ("Report", "Portfolio audit workbook"),
        ("Generated month", generated_local.strftime("%B %Y")),
        ("Generated date", generated_local.strftime("%Y-%m-%d")),
        ("Generated time", generated_local.strftime("%I:%M:%S %p")),
        ("Generated at", generated_local.strftime("%B %d, %Y %I:%M:%S %p")),
        ("Generated by", generated_by.display_name() if generated_by else ""),
        ("Period", cashflow["lookback_label"]),
        ("Date from", cashflow["date_from"].isoformat()),
        ("Date to", cashflow["date_to"].isoformat()),
        ("Group by", cashflow["granularity_label"]),
        ("Staff", cashflow["staff_label"]),
        ("Product", details["product_label"]),
        ("Classification", "Confidential — internal audit use"),
    ]


def _write_cover(wb, cashflow, details, generated_local, generated_by, org_name, tagline):
    ws = wb.active
    ws.title = "Cover"
    columns = 8
    _write_banner(
        ws,
        org_name,
        "Portfolio audit report",
        tagline or "Line-level collections, applications, disbursements, and transactions.",
        columns,
        tab_color=THEME_DEEP,
    )
    _apply_print(ws, "Cover", org_name)
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 42
    for letter, width in (("C", 16), ("D", 18), ("E", 16), ("F", 18), ("G", 16), ("H", 18)):
        ws.column_dimensions[letter].width = width

    ws.cell(5, 1, "Report parameters").font = _font(size=12, bold=True, color=THEME_DEEP)
    ws.merge_cells("A5:B5")
    meta_pairs = _meta_pairs(cashflow, details, generated_local, generated_by, org_name)
    meta_start = 6
    for index, (label, value) in enumerate(meta_pairs, start=meta_start):
        key = ws.cell(index, 1, label)
        key.fill = _fill(MINT)
        key.font = _font(size=10, bold=True, color=THEME_DEEP)
        key.border = THIN
        val = ws.cell(index, 2, value)
        val.border = THIN
        val.font = _font(size=10)

    kpi_row = 6
    kpis = [
        ("Collections", cashflow["collection"]["total_count"], cashflow["collection"]["total_amount"]),
        ("Applications", len(details["applications"]), sum((row["amount"] for row in details["applications"]), Decimal("0.00"))),
        ("Disbursements", len(details["disbursements"]), sum((row["principal"] for row in details["disbursements"]), Decimal("0.00"))),
        ("Savings", len(details["savings_transactions"]), cashflow["savings"]["total_amount"]),
        ("Mutual aid", cashflow["mutual_aid"]["total_count"], cashflow["mutual_aid"]["total_amount"]),
        ("All activity", len(details["ledger"]), None),
    ]
    ws.cell(5, 4, "Period snapshot").font = _font(size=12, bold=True, color=THEME_DEEP)
    ws.merge_cells("D5:F5")
    headers = ("Feature", "Count", "Amount")
    for col, title in enumerate(headers, start=4):
        cell = ws.cell(kpi_row, col, title)
        cell.fill = _fill(THEME)
        cell.font = _font(size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN
    for offset, (label, count, amount) in enumerate(kpis, start=1):
        ws.cell(kpi_row + offset, 4, label).border = THIN
        count_cell = ws.cell(kpi_row + offset, 5, count)
        count_cell.border = THIN
        count_cell.alignment = Alignment(horizontal="right")
        amount_cell = ws.cell(kpi_row + offset, 6, _as_number(amount) if amount is not None else "—")
        amount_cell.border = THIN
        if amount is not None:
            amount_cell.number_format = CURRENCY
        if offset % 2 == 0:
            for col in range(4, 7):
                ws.cell(kpi_row + offset, col).fill = _fill(ALT_ROW)

    contents_row = max(meta_start + len(meta_pairs) - 1, kpi_row + len(kpis)) + 2
    ws.cell(contents_row, 1, "Workbook tabs").font = _font(size=12, bold=True, color=THEME_DEEP)
    ws.merge_cells(start_row=contents_row, start_column=1, end_row=contents_row, end_column=2)
    contents = [
        ("Cover", "Report stamp, filters, and snapshot totals"),
        ("All Activity", "Chronological ledger of every recorded event"),
        ("Collections", "Loan payments, including savings and mutual-aid splits"),
        ("Applications", "Loan applications submitted in the period"),
        ("Disbursements", "Funds released to members"),
        ("Savings", "Deposits, interest, and withdrawals"),
        ("Mutual Aid", "Mutual aid contributions"),
        ("Period Totals", "Grouped totals matching the on-screen report"),
    ]
    head = contents_row + 1
    for col, title in enumerate(("Tab", "Contents"), start=1):
        cell = ws.cell(head, col, title)
        cell.fill = _fill(THEME)
        cell.font = _font(size=10, bold=True, color=WHITE)
        cell.border = THIN
    for offset, (sheet, description) in enumerate(contents, start=1):
        link = ws.cell(head + offset, 1, sheet)
        link.hyperlink = f"#'{sheet}'!A1"
        link.font = _font(size=10, bold=True, color=THEME)
        link.border = THIN
        desc = ws.cell(head + offset, 2, description)
        desc.border = THIN
        desc.font = _font(size=10)
        if offset % 2 == 0:
            link.fill = _fill(ALT_ROW)
            desc.fill = _fill(ALT_ROW)
    note = ws.cell(head + len(contents) + 2, 1, "Use the tabs at the bottom of this workbook to review each feature. Filters and freeze panes are enabled on every data sheet.")
    note.font = _font(size=9, color=MUTED)
    ws.merge_cells(start_row=head + len(contents) + 2, start_column=1, end_row=head + len(contents) + 2, end_column=6)
    return ws


def _write_period_totals(wb, cashflow, org_name, subtitle):
    ws = wb.create_sheet("Period Totals")
    columns = 6
    _write_banner(ws, org_name, "Period totals", subtitle, columns, tab_color=GOLD)
    _apply_print(ws, "Period totals", org_name)
    ws.freeze_panes = "A5"
    for index, width in enumerate((28, 16, 16, 16, 16, 16), start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    def write_block(start_row, title, headers, rows, total_row):
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=len(headers))
        heading = ws.cell(start_row, 1, title)
        heading.font = _font(size=12, bold=True, color=THEME_DEEP)
        heading.fill = _fill(MINT)
        for col in range(1, len(headers) + 1):
            ws.cell(start_row, col).fill = _fill(MINT)
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(start_row + 1, col, header)
            cell.fill = _fill(THEME)
            cell.font = _font(size=10, bold=True, color=WHITE)
            cell.border = THIN
            cell.alignment = Alignment(horizontal="center")
        for offset, row in enumerate(rows, start=1):
            for col, value in enumerate(row, start=1):
                cell = ws.cell(start_row + 1 + offset, col, _as_number(value) if col > 1 else value)
                cell.border = THIN
                if col > 1 and not isinstance(value, str):
                    if headers[col - 1].lower() in {"payments", "contributions", "deposits count", "count"}:
                        cell.number_format = "0"
                    else:
                        cell.number_format = CURRENCY
                if offset % 2 == 0:
                    cell.fill = _fill(ALT_ROW)
        total_idx = start_row + 2 + len(rows)
        for col, value in enumerate(total_row, start=1):
            cell = ws.cell(total_idx, col, _as_number(value) if col > 1 else value)
            cell.border = THIN
            cell.fill = _fill(MINT)
            cell.font = _font(size=10, bold=True, color=THEME_DEEP)
            if col > 1:
                cell.number_format = CURRENCY if headers[col - 1].lower() not in {"payments", "contributions", "deposits count", "count"} else "0"
        return total_idx + 2

    row = 5
    row = write_block(
        row,
        "Collections",
        ["Period", "Payments", "Amount"],
        [[item["label"], item["count"], item["amount"]] for item in cashflow["collection"]["rows"]],
        ["Total", cashflow["collection"]["total_count"], cashflow["collection"]["total_amount"]],
    )
    row = write_block(
        row,
        "Mutual aid",
        ["Period", "Contributions", "Amount"],
        [[item["label"], item["count"], item["amount"]] for item in cashflow["mutual_aid"]["rows"]],
        ["Total", cashflow["mutual_aid"]["total_count"], cashflow["mutual_aid"]["total_amount"]],
    )
    write_block(
        row,
        "Savings",
        ["Period", "Deposits count", "Deposits", "Interest", "Withdrawals", "Net"],
        [
            [item["label"], item["count"], item["amount"], item["interest"], item["withdrawals"], item["net"]]
            for item in cashflow["savings"]["rows"]
        ],
        [
            "Total",
            cashflow["savings"]["total_count"],
            cashflow["savings"]["total_amount"],
            cashflow["savings"]["total_interest"],
            cashflow["savings"]["total_withdrawals"],
            cashflow["savings"]["total_net"],
        ],
    )
    return ws


def build_audit_workbook(cashflow, *, generated_at, generated_by, product=None):
    details = audit_detail_context(
        cashflow["date_from"],
        cashflow["date_to"],
        staff_user=cashflow["selected_staff"],
        product=product,
    )
    generated_local = timezone.localtime(generated_at) if timezone.is_aware(generated_at) else generated_at
    features = Features.load()
    org_name = features.store_name or Features.DEFAULT_STORE_NAME
    tagline = features.tagline or Features.DEFAULT_TAGLINE
    subtitle = (
        f"{cashflow['lookback_label']}  ·  {cashflow['date_from']:%b %d, %Y} – {cashflow['date_to']:%b %d, %Y}"
        f"  ·  {cashflow['staff_label']}  ·  {details['product_label']}"
        f"  ·  Generated {generated_local.strftime('%B %d, %Y %I:%M:%S %p')}"
    )

    wb = Workbook()
    wb.properties.title = f"{org_name} portfolio audit"
    wb.properties.creator = generated_by.display_name() if generated_by else org_name
    wb.properties.subject = "Confidential internal audit workbook"
    wb.properties.description = subtitle

    _write_cover(wb, cashflow, details, generated_local, generated_by, org_name, tagline)
    _add_data_sheet(wb, "All Activity", "All activity", subtitle, org_name, LEDGER_COLUMNS, details["ledger"], tab_color=THEME)
    _add_data_sheet(wb, "Collections", "Collections", subtitle, org_name, COLLECTION_COLUMNS, details["collections"], tab_color="2E7D32")
    _add_data_sheet(wb, "Applications", "Applications", subtitle, org_name, APPLICATION_COLUMNS, details["applications"], tab_color="1565C0")
    _add_data_sheet(wb, "Disbursements", "Disbursements", subtitle, org_name, DISBURSEMENT_COLUMNS, details["disbursements"], tab_color=GOLD)
    _add_data_sheet(wb, "Savings", "Savings transactions", subtitle, org_name, SAVINGS_COLUMNS, details["savings_transactions"], tab_color="00838F")
    _add_data_sheet(wb, "Mutual Aid", "Mutual aid contributions", subtitle, org_name, MUTUAL_AID_COLUMNS, details["mutual_aid_contributions"], tab_color="6D4C41")
    _write_period_totals(wb, cashflow, org_name, subtitle)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
